from openpyxl import load_workbook


class ExcelUtils:

    @staticmethod
    def read_records(file_path):

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[2]
        ]

        try:
            order_no_index = headers.index("Order No")
        except ValueError:
            workbook.close()
            raise AssertionError(
                f"'Order No' column not found. "
                f"Available headers: {headers}"
            )

        records = []

        for row in worksheet.iter_rows(
            min_row=3,
            values_only=True
        ):
            order_no = row[order_no_index]

            if order_no is None:
                continue

            if str(order_no).strip() == "":
                continue

            records.append(
                dict(zip(headers, row))
            )

        workbook.close()

        return records

    @staticmethod
    def get_row_count(file_path):

        records = ExcelUtils.read_records(file_path)

        return len(records)

    @staticmethod
    def get_data_row_count(file_path):

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        worksheet = workbook.active

        data_row_count = 0

        for row in worksheet.iter_rows(
            min_row=3,
            values_only=True
        ):
            if any(
                value is not None and str(value).strip() != ""
                for value in row
            ):
                data_row_count += 1

        workbook.close()

        return data_row_count